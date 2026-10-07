# -*- coding: utf-8 -*-
"""Planilha da conferência de lojas (Aldo 06/10/2026). Lê o -calc.json e grava docs/Conferencia-lojas-AIVA-2026-10-06.xlsx."""
import json, io
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

d = json.load(io.open('scripts/out-conferencia-lojas-2026-10-06-calc.json', encoding='utf-8'))
L = d['linhas']; FORA = d['fora']

GRUPOS = {
    '1': ('OK — comissionada', 'Loja nossa, criada na Track e com comissão nos relatórios de jul/ago.', 'Nada.'),
    '2': ('Vendeu e falta comissão', 'Tem venda no portal da Track num mês com relatório de comissão, mas não aparece comissão nesse mês.', 'Cobrar a comissão da AIVA (Mauricio).'),
    '3': ('Loja da Track — vendas só em set/out', 'Loja nossa na Track; as vendas são de setembro/outubro e o relatório de comissão desses meses ainda não chegou.', 'Conferir quando chegar o relatório de setembro.'),
    '4': ('Loja da Track sem venda', 'Loja criada na Track (tem RID) e sem nenhuma venda desde maio.', 'Ativação: ligar / treinar.'),
    '5': ('SUSPEITA — concluída por nós e fora da Track', 'Nós levamos até o fim (status/card em Treinar-Login-Vendendo ou funil 11 AIVA), mas o CNPJ NÃO está na Track no portal. Quando o CNPJ está na base de clientes AIVA, a loja opera por OUTRO canal/parceiro.', 'Mandar a lista pro Mauricio confirmar em nome de quem a loja está.'),
    '6': ('Travada antes de virar loja', 'Passou da pré-aprovação, mas não virou loja: pré-cadastro não chegou, formulário/biometria pendente ou reprovada.', 'Funil (cobrança/biometria já automáticas).'),
    '7': ('Concluída por nós SEM CNPJ no cadastro', 'Lead fechado (Treinar/Login/Vendendo ou funil 11) sem nenhum CNPJ gravado do nosso lado — não dá pra cruzar. A coluna "Possível loja pelo nome" sugere o CNPJ.', 'Completar o CNPJ no cadastro e reconferir.'),
    '8': ('Na Track com OUTRO CNPJ (pelo telefone)', 'O CNPJ que temos não está na Track, mas o telefone do lead casa com uma loja da Track de outro CNPJ (filial, outra empresa do sócio, digitação).', 'Normalmente OK; conferir se o CNPJ certo está no nosso cadastro.'),
    '9': ('Sem CNPJ (não concluída)', 'Passou da pré-aprovação sem CNPJ gravado e não chegou ao fim.', 'Nada (funil).'),
}
COLS = [
    ('grupo', 'Grupo', 30), ('motivo', 'Por quê / sinais', 55), ('cnpj', 'CNPJ', 17), ('loja', 'Loja', 32), ('telefone', 'Telefone', 15),
    ('status_nosso', 'Status (nosso)', 24), ('etapa_evo', 'Etapa no Evo', 16), ('entrada', 'Entrada (pré-aprov.)', 12), ('funil11', 'Funil 11 (MRR)', 12),
    ('onb_stage', 'Onboarding AIVA', 18), ('biometria', 'Biometria', 11), ('rid', 'RID (Track)', 10), ('status_portal', 'Status portal', 11), ('bloqueada', 'Bloq. limite', 9),
    ('consultas', 'Consultas', 10), ('vendas', 'Vendas (qtd)', 10), ('valor_vendas', 'Vendido R$ (mai-out)', 14), ('vendas_jul', 'Vendido jul', 12), ('com_jul', 'Comissão jul', 12),
    ('vendas_ago', 'Vendido ago', 12), ('com_ago', 'Comissão ago', 12), ('vendas_set_out', 'Vendido set+out', 13), ('comissao_total', 'Comissão total', 12),
    ('meses_sem_comissao', 'Meses c/ venda e sem comissão', 14), ('cnpj_track_pelo_tel', 'CNPJ na Track pelo telefone', 20), ('match_nome', 'Possível loja pelo nome (conferir)', 55),
    ('base_aiva', 'Na base de clientes AIVA como', 40),
]
MOEDA = {'valor_vendas', 'vendas_jul', 'com_jul', 'vendas_ago', 'com_ago', 'vendas_set_out', 'comissao_total'}
TEXTO = {'cnpj', 'telefone', 'rid', 'cnpj_track_pelo_tel'}
F = Font(name='Arial', size=10); FB = Font(name='Arial', size=10, bold=True, color='FFFFFF'); FT = Font(name='Arial', size=14, bold=True)
HEAD = PatternFill('solid', fgColor='1F4E78')
COR = {'1': 'E2EFDA', '2': 'FCE4D6', '3': 'DDEBF7', '4': 'FFF2CC', '5': 'F8CBAD', '6': 'EDEDED', '7': 'FFE699', '8': 'E2EFDA', '9': 'F2F2F2'}

wb = Workbook()
ws = wb.active; ws.title = 'Resumo'
ws['A1'] = 'Conferência das lojas AIVA — nosso funil × portal da AIVA × comissão'; ws['A1'].font = FT
ws['A2'] = 'Universo: todo lead que chegou na pré-aprovação desde maio/2026 (+ funil 11 com etiqueta AIVA), uma linha por CNPJ. Gerado em 06/10/2026.'; ws['A2'].font = F
ws['A3'] = 'Comissão: relatórios de julho e agosto/2026 (carteira + FCDL). Vendas de setembro/outubro ainda sem relatório.'; ws['A3'].font = F
hdr = ['Grupo', 'Lojas (CNPJs)', 'Vendido R$ (mai-out)', 'Comissão R$ (jul+ago)', 'O que é', 'O que fazer']
for i, h in enumerate(hdr, 1):
    c = ws.cell(row=5, column=i, value=h); c.font = FB; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical='center')
r = 6
for g, (nome, oque, acao) in GRUPOS.items():
    rows = [l for l in L if l['grupo'].startswith(g)]
    vals = [f'{g}. {nome}', len(rows), round(sum(l['valor_vendas'] for l in rows), 2), round(sum(l['comissao_total'] for l in rows), 2), oque, acao]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v); c.font = F; c.alignment = Alignment(wrap_text=True, vertical='top')
        c.fill = PatternFill('solid', fgColor=COR[g])
        if i in (3, 4): c.number_format = 'R$ #,##0.00'
    r += 1
ws.cell(row=r, column=1, value='Total').font = Font(name='Arial', size=10, bold=True)
ws.cell(row=r, column=2, value=len(L)).font = Font(name='Arial', size=10, bold=True)
r += 2
ws.cell(row=r, column=1, value=f'Lojas da Track no portal que NÃO estão no nosso cadastro: {len(FORA)} (aba "Fora do nosso cadastro").').font = F
r += 2
# entrada por mês × grupo principal
ws.cell(row=r, column=1, value='Por mês de entrada (pré-aprovação)').font = Font(name='Arial', size=11, bold=True); r += 1
meses = sorted({l['entrada'][:7] for l in L if l['entrada']})
cab = ['Mês'] + [f'{g}.' for g in GRUPOS] + ['Total']
for i, h in enumerate(cab, 1):
    c = ws.cell(row=r, column=i, value=h); c.font = FB; c.fill = HEAD
r += 1
for m in meses:
    rows = [l for l in L if l['entrada'][:7] == m]
    vals = [m] + [sum(1 for l in rows if l['grupo'].startswith(g)) for g in GRUPOS] + [len(rows)]
    for i, v in enumerate(vals, 1): ws.cell(row=r, column=i, value=v).font = F
    r += 1
for i, w in enumerate([44, 13, 18, 18, 70, 45], 1): ws.column_dimensions[get_column_letter(i)].width = w

def aba(titulo, linhas, cols=COLS):
    w = wb.create_sheet(titulo[:31])
    for i, (k, h, wd) in enumerate(cols, 1):
        c = w.cell(row=1, column=i, value=h); c.font = FB; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical='center')
        w.column_dimensions[get_column_letter(i)].width = wd
    for ri, l in enumerate(linhas, 2):
        for i, (k, h, wd) in enumerate(cols, 1):
            v = l.get(k, '')
            c = w.cell(row=ri, column=i, value=(str(v) if k in TEXTO and v not in ('', None) else v)); c.font = F
            if k in TEXTO: c.number_format = '@'
            if k in MOEDA: c.number_format = 'R$ #,##0.00'
    w.freeze_panes = 'D2'
    w.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{max(len(linhas) + 1, 1)}'
    w.row_dimensions[1].height = 32

ordem = lambda l: (l['grupo'], -l['valor_vendas'], l['loja'])
aba('Para o Mauricio (suspeitas)', sorted([l for l in L if l['grupo'][0] in '52'], key=ordem))
for g, (nome, _, _) in GRUPOS.items():
    rows = sorted([l for l in L if l['grupo'].startswith(g)], key=ordem)
    if rows: aba(f'{g}. {nome}'.replace('/', '-').replace('—', '-'), rows)
aba('Todas', sorted(L, key=ordem))
aba('Fora do nosso cadastro', sorted(FORA, key=lambda x: -x['valor_vendas']), [
    ('cnpj', 'CNPJ', 17), ('loja', 'Loja (portal)', 50), ('rid', 'RID', 10), ('cadastro', 'Cadastro no portal', 14), ('status_portal', 'Status', 10),
    ('valor_vendas', 'Vendido R$ (mai-out)', 15), ('comissao_total', 'Comissão total', 13), ('na_onboarding', 'No onboarding', 12), ('funil11', 'Funil 11', 10)])
wb.save('docs/Conferencia-lojas-AIVA-2026-10-06.xlsx')
print('ok')
