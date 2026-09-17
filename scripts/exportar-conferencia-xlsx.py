# -*- coding: utf-8 -*-
"""
Monta o xlsx da conferência mensal de comissão UME a partir do JSON gerado por
scripts/exportar-conferencia-json.mjs (+ JSON do mês anterior em ume_comissoes).

Uso:
  python scripts/exportar-conferencia-xlsx.py --conf conf.json --anterior ume-anterior.json --out arquivo.xlsx
"""
import argparse, json, re, unicodedata, collections, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ap = argparse.ArgumentParser()
ap.add_argument("--conf", required=True)
ap.add_argument("--anterior", required=True)
ap.add_argument("--out", required=True)
ap.add_argument("--nf-total", type=float, default=None)
ap.add_argument("--nf-odres", type=float, default=None)
ap.add_argument("--nf-ume", type=float, default=None)
a = ap.parse_args()

d = json.load(open(a.conf, encoding="utf-8"))
conf = d["conf"]
des = {x["cnpj"]: x for x in d["desempenho"]}
ant = json.load(open(a.anterior, encoding="utf-8"))
mes = d["mes"]

def n(s):
    return re.sub(r"[^a-z0-9]", "", unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower())

H = Font(name="Arial", bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F3864"); B = Font(name="Arial", size=10)
wb = openpyxl.Workbook(); wb.remove(wb.active)

def sheet(title, hdr, rows, text_cols=()):
    ws = wb.create_sheet(title); ws.append(hdr)
    for c in ws[1]:
        c.font = H; c.fill = HF
    for r in rows:
        ws.append(r)
    for j, h in enumerate(hdr, 1):
        L = get_column_letter(j)
        w = max([len(str(h))] + [len(str(r[j - 1] or "")) for r in rows])
        ws.column_dimensions[L].width = min(max(w + 2, 10), 55)
        for c in ws[L][1:]:
            c.font = B
            if h in text_cols:
                c.number_format = "@"
            elif isinstance(c.value, float):
                c.number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws

rel = [c["relatorio"] for c in conf if c["relatorio"]]
a_rid = {r["retailer_id"] for r in rel if r["retailer_id"]}
a_cnpj = {r["cnpj"] for r in rel if r["cnpj"]}
sumiram = [r for r in ant if r["retailer_id"] not in a_rid and r["cnpj"] not in a_cnpj]
sum_cart = [r for r in sumiram if r["origem"] == "carteira"]
est = collections.Counter(c["estado"] for c in conf)
so_cart = [c["relatorio"] for c in conf if c["estado"] == "so_relatorio" and c["relatorio"]["origem"] == "carteira"]
portal_fora = [x for x in d["desempenho"] if (x["vendas"] or 0) > 0 and x["cnpj"] not in a_cnpj and str(x.get("rid") or "") not in {str(r) for r in a_rid}]
byc = {r["cnpj"]: r for r in rel if r["cnpj"]}
qtd_dif = [x for x in d["desempenho"] if x["cnpj"] in byc and (x["vendas"] or 0) != (byc[x["cnpj"]]["contratos"] or 0)]
cn = collections.defaultdict(list)
for c in conf:
    if c["opp"] and c["cnpj"]:
        cn[c["cnpj"]].append(c)
dups = {k: v for k, v in cn.items() if len(v) > 1}
falsos = [c for c in conf if c["divergencia"] and c["cnpj"] in dups]
reais = [c for c in conf if c["divergencia"] and c["cnpj"] not in dups]

# ── Resumo
res = wb.create_sheet("Resumo")
linhas = [(f"Conferência de comissão Ume — {mes}", ""), ("Gerado por", "scripts/importar-comissao-ume.mjs → exportar-conferencia-json.mjs → exportar-conferencia-xlsx.py"), ("", "")]
if a.nf_total:
    linhas += [("Valor total da NF (Carteira + FCDL)", a.nf_total)]
    if a.nf_odres: linhas += [("  Nota 1 · Odres", a.nf_odres)]
    if a.nf_ume: linhas += [("  Nota 2 · Ume", a.nf_ume)]
    linhas += [("", "")]
linhas += [
    ("RESULTADO — lojas com venda no Portal AIVA e AUSENTES do relatório", len(portal_fora)),
    ("RESULTADO — lojas casadas com qtd de contratos diferente do Portal", len(qtd_dif)),
    ("RESULTADO — divergências reais do painel (sem venda no relatório, venda no portal)", len(reais)),
    ("Divergências FALSAS do painel (conta duplicada no funil 11)", len(falsos)),
    ("", ""),
    ("Contas do funil 11 (UME/AIVA) conferidas", d["contasUmeAiva"]),
    ("  Comissionadas", est["comissionada"]),
    ("  Sem venda no mês (com Retailer ID)", est["sem_venda"]),
    ("  Sem Retailer ID na conta", est["sem_rid"]),
    ("Linhas do relatório sem conta no funil 11 — carteira (pagas, não rastreadas)", len(so_cart)),
    ("  Comissão dessas linhas", round(sum(r["comissao"] or 0 for r in so_cart), 2)),
    ("Linhas do relatório sem conta no funil 11 — FCDL (esperado)", est["so_relatorio"] - len(so_cart)),
    ("", ""),
    ("Lojas comissionadas no MÊS ANTERIOR e fora deste mês (carteira)", len(sum_cart)),
    ("  Comissão do mês anterior dessas lojas", round(sum(r["comissao"] or 0 for r in sum_cart), 2)),
    ("CNPJs duplicados no funil 11 (espelho AIVA × conta UME)", len(dups)),
]
for r in linhas:
    res.append(r)
res.column_dimensions["A"].width = 80; res.column_dimensions["B"].width = 60
for row in res.iter_rows():
    for c in row:
        c.font = B; c.alignment = Alignment(wrap_text=True, vertical="top")
res["A1"].font = Font(name="Arial", size=13, bold=True)
for row in res.iter_rows():
    if str(row[0].value or "").startswith("RESULTADO"):
        row[0].font = Font(name="Arial", size=10, bold=True)

# ── Portal com venda fora do relatório (o que importa)
sheet("Portal vende, Ume nao paga", ["RID", "CNPJ", "Loja (portal)", "Vendas", "Valor vendas", "Status portal"],
      [[x.get("rid"), x["cnpj"], x["loja"], x["vendas"], x["valor_vendas"], x.get("status_portal")] for x in portal_fora], text_cols=("CNPJ",))

# ── Divergências do painel (reais e falsas)
rows = []
for c in reais + falsos:
    p = des.get(c["cnpj"]) or {}
    rows.append(["REAL" if c in reais else "FALSA (duplicata)", c["opp"]["id"], c["opp"]["title"], c["umeRid"] or "", c["cnpj"], p.get("vendas"), p.get("valor_vendas"),
                 "; ".join(f'opp {o["opp"]["id"]} {o["opp"]["title"][:30]} → {o["estado"]}' for o in dups.get(c["cnpj"], []) if o is not c)])
sheet("Divergencias painel", ["Tipo", "Opp ID", "Loja (funil)", "UME_RID", "CNPJ", "Vendas ago (portal)", "Valor", "Outras contas com o mesmo CNPJ"], rows, text_cols=("CNPJ",))

# ── Mês anterior sim, este mês não
rows = []
for r in sorted(sumiram, key=lambda r: (r["origem"] != "carteira", -(r["contratos"] or 0))):
    p = des.get(r["cnpj"]) or {}
    leitura = "Portal AIVA confirma 0 venda" if p and (p.get("vendas") or 0) == 0 else ("VERIFICAR — portal mostra venda" if p else "sem contraprova (loja não-AIVA)")
    rows.append([r["origem"], r["retailer_id"], r["cnpj"], r["varejo"], r["grupo"], r["contratos"], r["comissao"], p.get("status_portal") or "", p.get("aprovados"), p.get("vendas"), leitura])
sheet("Mes anterior sim, este nao", ["Origem", "Retailer ID", "CNPJ", "Varejo", "Grupo", "Contratos (mês ant.)", "Comissão (mês ant.)", "Status Portal AIVA", "Aprovados (portal)", "Vendas (portal)", "Leitura"], rows, text_cols=("CNPJ",))

# ── Pagas fora do funil
rows = [[r["retailer_id"], r["cnpj"], r["varejo"], r["grupo"], r["contratos"], r["originacao"], r["mdr"], r["comissao"], (des.get(r["cnpj"]) or {}).get("status_portal") or ""]
        for r in sorted(so_cart, key=lambda r: -(r["comissao"] or 0))]
sheet("Pagas fora do funil (cadastrar)", ["Retailer ID", "CNPJ", "Varejo", "Grupo", "Contratos", "Originação", "MDR", "Comissão", "Status Portal AIVA"], rows, text_cols=("CNPJ",))

# ── Sem Retailer ID
so = [c["relatorio"] for c in conf if c["estado"] == "so_relatorio"]
rows = []
for c in [c for c in conf if c["estado"] == "sem_rid"]:
    t = n(c["opp"]["title"]); cand = ""
    for r in so:
        v = n(r["varejo"])
        if len(t) >= 8 and (t in v or v in t):
            cand = f'{r["retailer_id"]} · {r["varejo"]} · {r["grupo"]} · R$ {r["comissao"]:.2f}'; break
    rows.append([c["opp"]["id"], c["opp"]["title"], "AIVA" if 69 in (c["tags"] or []) else "UME", c["cnpj"] or "", cand, (des.get(c["cnpj"]) or {}).get("vendas") if c["cnpj"] else ""])
sheet("Sem Retailer ID", ["Opp ID", "Loja (funil)", "Etiqueta", "CNPJ na descrição", "Candidato no relatório (por nome)", "Vendas (portal)"], rows, text_cols=("CNPJ na descrição",))

# ── Duplicatas
rows = []
for k, v in dups.items():
    for c in v:
        rows.append([k, c["opp"]["id"], c["opp"]["title"], "AIVA" if 69 in (c["tags"] or []) else "UME", c["umeRid"] or "", c["estado"], (c["relatorio"] or {}).get("comissao")])
sheet("Duplicatas funil 11", ["CNPJ", "Opp ID", "Loja (funil)", "Etiqueta", "UME_RID", "Estado", "Comissão"], rows, text_cols=("CNPJ",))

# ── Detalhe
rows = []
for c in conf:
    if not c["opp"]:
        continue
    r = c["relatorio"] or {}; p = des.get(c["cnpj"]) or {}
    rows.append([c["estado"], c["opp"]["id"], c["opp"]["title"], "AIVA" if 69 in (c["tags"] or []) else "UME", c["umeRid"] or "", c["cnpj"] or "", r.get("retailer_id"), r.get("varejo"), r.get("grupo"), r.get("contratos"), r.get("comissao"), p.get("vendas"), p.get("status_portal")])
rows.sort(key=lambda r: (r[0], -(r[10] or 0)))
sheet("Detalhe funil 11", ["Estado", "Opp ID", "Loja (funil)", "Etiqueta", "UME_RID", "CNPJ", "RID relatório", "Varejo relatório", "Grupo", "Contratos", "Comissão", "Vendas (portal)", "Status portal"], rows, text_cols=("CNPJ",))

os.makedirs(os.path.dirname(a.out), exist_ok=True)
wb.save(a.out)
print("OK", a.out)
print(f"portal_fora={len(portal_fora)} qtd_dif={len(qtd_dif)} reais={len(reais)} falsas={len(falsos)} sumiram_carteira={len(sum_cart)} dups={len(dups)}")
