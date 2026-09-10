# -*- coding: utf-8 -*-
"""Excel editável do fluxo da VictorIA — mesma fonte de dados do gerador do mapa."""
import io, sys, re
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── carrega COLS / SAIDAS / TRANSVERSAL / LANES / movers do gerador (só a parte de dados)
src = io.open('scripts/gerar-fluxo-victoria.py', encoding='utf-8').read()
dados = src.split('# ─────────────────────────── DRAW.IO')[0]
ns = {}
exec(dados, ns)
COLS, SAIDAS, TRANSVERSAL, LANES = ns['COLS'], ns['SAIDAS'], ns['TRANSVERSAL'], ns['LANES']
MOVERS = ['VictorIA (lead conversa)', 'VictorIA (7 dados)', 'Nei, após OK do Edu', 'Nei (12 dados)', 'Nei, após OK do Edu', 'Nei (treinou)', 'Nei (1ª venda)']

LANE_NOME = dict(victoria='VictorIA', humano='Humanos', rotinas='Rotinas automáticas', sistemas='Sistemas e dados')
COR = dict(evo=('DCE8F7', '1F4E8C'), victoria=('EBE5F9', '5B3FA6'), humano=('FBEBC8', '8A5200'),
           rotinas=('D6F0EC', '0F6E63'), sistemas=('E7EAEE', '3E4A57'), saida=('F6DEDA', '8C2F2F'), trans=('F1EFE8', '4F4A3A'))
ARIAL = 'Arial'
thin = Side(style='thin', color='C9CED6')
BORDA = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')

def fonte(bold=False, color='1B2430', size=10):
    return Font(name=ARIAL, bold=bold, color=color, size=size)
def fill(hexa):
    return PatternFill('solid', fgColor=hexa)
def cabecalho(ws, row, valores, cor_fundo='1F4E8C', cor_txt='FFFFFF'):
    for i, v in enumerate(valores, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = fonte(True, cor_txt); c.fill = fill(cor_fundo); c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = BORDA
def largura(ws, larguras):
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

wb = Workbook()

# ── Aba 1: Resumo
ws = wb.active; ws.title = 'Resumo'
ws['A1'] = 'Fluxo atual da VictorIA — AIVA'; ws['A1'].font = fonte(True, size=14)
ws['A2'] = 'Retrato de 10/09/2026 tirado do código em produção. Colunas do mapa = etapas do funil 15 do Evo Talks; linhas = quem age. Edite à vontade — a aba "Itens" tem a coluna NOVO FLUXO em branco pra você redesenhar.'
ws['A2'].font = fonte(); ws['A2'].alignment = WRAP; ws.merge_cells('A2:G2'); ws.row_dimensions[2].height = 42
cabecalho(ws, 4, ['Ordem', 'Etapa no Evo', 'ID', 'Automação de entrada', 'Como o card chega', 'Quem move pra próxima', 'Itens VictorIA', 'Itens Humanos', 'Itens Rotinas', 'Itens Sistemas'])
for i, c in enumerate(COLS):
    r = 5 + i
    vals = [i + 1, c['nome'], c['num'], c['auto'], c['chega'], MOVERS[i] if i < len(MOVERS) else '— (etapa final)',
            len(c['victoria']), len(c['humano']), len(c['rotinas']), len(c['sistemas'])]
    for j, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=j, value=v); cell.font = fonte(); cell.alignment = WRAP; cell.border = BORDA
        if j == 2: cell.font = fonte(True); cell.fill = fill(COR['evo'][0])
r = 5 + len(COLS) + 1
ws.cell(row=r, column=1, value='Saídas do funil').font = fonte(True, size=11)
cabecalho(ws, r + 1, ['', 'Saída', 'ID', 'Como entra', 'O que acontece'], cor_fundo='8C2F2F')
for k, s in enumerate(SAIDAS['itens']):
    rr = r + 2 + k
    for j, v in enumerate(['', s['nome'], s['num'] or '', s['auto'], s['txt']], 1):
        cell = ws.cell(row=rr, column=j, value=v); cell.font = fonte(); cell.alignment = WRAP; cell.border = BORDA
        if j == 2: cell.fill = fill(COR['saida'][0]); cell.font = fonte(True)
largura(ws, [7, 28, 6, 20, 34, 24, 10, 10, 10, 10])
ws.freeze_panes = 'A5'

# ── Aba 2: Fluxo (matriz) — igual à página
ws = wb.create_sheet('Fluxo (matriz)')
cab = ['Linha'] + [f'{c["num"]} · {c["nome"]}' for c in COLS] + ['Saídas do funil']
cabecalho(ws, 1, cab)
ws.row_dimensions[1].height = 32
# linha Etapa
ws.cell(row=2, column=1, value='Etapa no Evo\n(automação → como chega)')
for i, c in enumerate(COLS):
    ws.cell(row=2, column=2 + i, value=f'{c["auto"]} → opportunity-stage\n{c["chega"]}\n\nMove pra próxima: {MOVERS[i] if i < len(MOVERS) else "— etapa final"}')
ws.cell(row=2, column=2 + len(COLS), value='\n'.join(f'• {(str(s["num"])+" · ") if s["num"] else ""}{s["nome"]} ({s["auto"]}): {s["txt"]}' for s in SAIDAS['itens']))
for j in range(1, len(cab) + 1):
    cell = ws.cell(row=2, column=j); cell.font = fonte(bold=(j == 1), color=COR['evo'][1]); cell.fill = fill(COR['evo'][0] if j <= len(cab) - 1 else COR['saida'][0]); cell.alignment = WRAP; cell.border = BORDA
ws.row_dimensions[2].height = 150
for li, (lk, ltitle) in enumerate(LANES):
    r = 3 + li
    ws.cell(row=r, column=1, value=LANE_NOME[lk])
    maxlen = 0
    for i, c in enumerate(COLS):
        txt = '\n'.join(f'• {x}' for x in c[lk]); maxlen = max(maxlen, len(txt))
        ws.cell(row=r, column=2 + i, value=txt)
    ws.cell(row=r, column=2 + len(COLS), value='')
    for j in range(1, len(cab) + 1):
        cell = ws.cell(row=r, column=j); cell.font = fonte(bold=(j == 1), color=COR[lk][1] if j == 1 else '1B2430'); cell.fill = fill(COR[lk][0]) if j == 1 else fill('FFFFFF'); cell.alignment = WRAP; cell.border = BORDA
    ws.row_dimensions[r].height = min(400, 30 + maxlen // 2.2)
r = 3 + len(LANES)
ws.cell(row=r, column=1, value='Transversal\n(corre o funil inteiro)')
ws.cell(row=r, column=2, value='\n'.join(f'• {t}: {d}' for t, d in TRANSVERSAL))
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=len(cab))
for j in (1, 2):
    cell = ws.cell(row=r, column=j); cell.font = fonte(bold=(j == 1), color=COR['trans'][1] if j == 1 else '1B2430'); cell.fill = fill(COR['trans'][0]) if j == 1 else fill('FFFFFF'); cell.alignment = WRAP; cell.border = BORDA
ws.row_dimensions[r].height = 190
largura(ws, [22] + [42] * len(COLS) + [46])
ws.freeze_panes = 'B2'

# ── Aba 3: Itens (lista editável)
ws = wb.create_sheet('Itens')
cabecalho(ws, 1, ['Ordem', 'Etapa (Evo)', 'ID etapa', 'Linha', 'Nº', 'Como é hoje', 'Manter?', 'Novo fluxo (editar)', 'Observação'])
r = 2
for i, c in enumerate(COLS):
    for lk, _ in LANES:
        for n, item in enumerate(c[lk], 1):
            vals = [i + 1, c['nome'], c['num'], LANE_NOME[lk], n, item, '', '', '']
            for j, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=j, value=v); cell.font = fonte(); cell.alignment = WRAP; cell.border = BORDA
                if j == 4: cell.fill = fill(COR[lk][0])
            r += 1
for k, s in enumerate(SAIDAS['itens']):
    vals = [99, 'Saídas do funil', s['num'] or '', 'Saída', k + 1, f'{s["nome"]} ({s["auto"]}): {s["txt"]}', '', '', '']
    for j, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=j, value=v); cell.font = fonte(); cell.alignment = WRAP; cell.border = BORDA
        if j == 4: cell.fill = fill(COR['saida'][0])
    r += 1
for k, (t, d) in enumerate(TRANSVERSAL):
    vals = [100, 'Transversal', '', 'Transversal', k + 1, f'{t}: {d}', '', '', '']
    for j, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=j, value=v); cell.font = fonte(); cell.alignment = WRAP; cell.border = BORDA
        if j == 4: cell.fill = fill(COR['trans'][0])
    r += 1
largura(ws, [7, 22, 8, 18, 5, 70, 9, 60, 30])
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:I{r - 1}'
# validação simples na coluna Manter?
from openpyxl.worksheet.datavalidation import DataValidation
dv = DataValidation(type='list', formula1='"sim,não,mudar"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f'G2:G{r - 1}')

# ── Aba 4: Etapas do Evo (referência)
ws = wb.create_sheet('Etapas Evo')
cabecalho(ws, 1, ['Ordem', 'ID', 'Etapa', 'Automação de entrada', 'Como o card chega', 'Quem move pra próxima', 'Status no painel'])
STATUS = {66: 'INICIO', 47: 'INTERESSADO', 54: 'PRE_APROVACAO', 49: 'CADASTRO_RECEBIDO (ou INTERESSADO enquanto a Fase 3 não fecha)', 50: 'EM_ANALISE_AIVA', 70: 'TREINAR', 71: 'LOGIN', 51: 'LOJA_FINALIZADA_E_VENDENDO',
          53: 'SEM_RESPOSTA', 69: 'BOT_DETECTADO', 93: 'NAO_QUALIFICADO', 19: 'ODRES / UME (lead sai do funil 15)'}
r = 2
for i, c in enumerate(COLS):
    for j, v in enumerate([i + 1, c['num'], c['nome'], c['auto'], c['chega'], MOVERS[i] if i < len(MOVERS) else '— etapa final', STATUS.get(c['num'], '')], 1):
        cell = ws.cell(row=r, column=j, value=v); cell.font = fonte(); cell.alignment = WRAP; cell.border = BORDA
    r += 1
for s in SAIDAS['itens']:
    for j, v in enumerate(['saída', s['num'] or '', s['nome'], s['auto'], s['txt'], '', STATUS.get(s['num'], '')], 1):
        cell = ws.cell(row=r, column=j, value=v); cell.font = fonte(); cell.alignment = WRAP; cell.border = BORDA; cell.fill = fill(COR['saida'][0]) if j == 3 else fill('FFFFFF')
    r += 1
largura(ws, [7, 6, 30, 20, 40, 24, 40])
ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:G{r - 1}'

out = sys.argv[1]
wb.save(out)
print('xlsx ok', out)
