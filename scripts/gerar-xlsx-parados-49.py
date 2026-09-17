# -*- coding: utf-8 -*-
"""Planilha do levantamento dos parados em Cadastro Recebido (etapa 49).

Lê docs/parados-49.json (gerado por scripts/parados-cadastro-recebido-2026-09-17.mjs)
e gera docs/Parados-Cadastro-Recebido-2026-09-17.xlsx.
"""
import io, json, os, sys
sys.stdout.reconfigure(encoding='utf-8')
os.chdir(r'C:\projetos claude\sdr-aiva')

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

linhas = json.load(io.open('docs/parados-49.json', encoding='utf-8'))

ARIAL = 'Arial'
HDR_BG = PatternFill('solid', fgColor='1F3864')
HDR_FT = Font(name=ARIAL, size=10, bold=True, color='FFFFFF')
CEL = Font(name=ARIAL, size=10)
BOLD = Font(name=ARIAL, size=10, bold=True)
TIT = Font(name=ARIAL, size=13, bold=True, color='1F3864')
BORDA = Border(*[Side(style='thin', color='D9D9D9')] * 4)
VERM = PatternFill('solid', fgColor='FCE4E4')
AMAR = PatternFill('solid', fgColor='FFF2CC')

wb = Workbook()

# ─── Resumo ──────────────────────────────────────────────────────────────────
ws = wb.active
ws.title = 'Resumo'
ws['A1'] = 'Oportunidades paradas em "Cadastro Recebido" (etapa 49) — 17/09/2026'
ws['A1'].font = TIT
ws['A2'] = f'{len(linhas)} cards no funil 15. "Parado" = dias desde que o card entrou na etapa (stagebegintime do Evo).'
ws['A2'].font = CEL
ws['A3'] = 'Um card sai da 49 quando: o lojista fecha a Fase 3 (e-mail) → o Nei lança o CNPJ na AIVA → a AIVA aprova → o espelho move pra 50.'
ws['A3'].font = CEL

por_trava = {}
for l in linhas:
    por_trava.setdefault(l['trava'], []).append(l)

ws['A5'] = 'O QUE ESTÁ SEGURANDO'
ws['A5'].font = BOLD
hdr = ['Trava', 'Cards', 'Mais antigo', 'De quem é a ação', 'O que fazer']
for i, h in enumerate(hdr, 1):
    c = ws.cell(row=6, column=i, value=h); c.fill = HDR_BG; c.font = HDR_FT; c.border = BORDA

DONO = {
    '2. Falta o e-mail (Fase 3 aberta)': ('Ninguém — régua encerrada', 'A automação já deu os 3 toques e escalou. Decisão do Aldo: descartar, ou fazer uma tentativa humana.'),
    '3. Sem CNPJ no painel': ('NOSSA (falha de sistema)', 'O CNPJ está nas observações do lead mas nunca virou linha no /registros. Criar o registro → o Nei lança.'),
    '4. CNPJ não lançado na AIVA': ('NEI', 'Lançar no form de pré-cadastro pelo /registros.'),
    '5. Lançado, não apareceu no portal': ('AIVA', 'Conferir com a AIVA por que não registrou.'),
    '6. Reprovado pela AIVA': ('NOSSA (bug)', 'O espelho deveria ter movido pra 95.'),
    '7. No portal, aguardando a AIVA': ('AIVA', 'O espelho move sozinho quando avançar.'),
    '1. Sem lead nosso': ('Conferir', 'Card sem lead correspondente.'),
}
r = 7
for trava in sorted(por_trava):
    grupo = por_trava[trava]
    dono, acao = DONO.get(trava, ('—', '—'))
    vals = [trava, len(grupo), f"{max(x['paradoDias'] or 0 for x in grupo)} dias", dono, acao]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v); c.font = CEL; c.border = BORDA
        c.alignment = Alignment(vertical='top', wrap_text=(i == 5))
        if 'NOSSA' in str(vals[3]): c.fill = VERM
    r += 1

r += 1
ws.cell(row=r, column=1, value='TEMPO PARADO').font = BOLD
r += 1
for rot, n in [('mais de 30 dias', 30), ('mais de 14 dias', 14), ('mais de 7 dias', 7)]:
    q = len([x for x in linhas if (x['paradoDias'] or 0) >= n])
    ws.cell(row=r, column=1, value=rot).font = CEL
    ws.cell(row=r, column=2, value=q).font = BOLD
    r += 1

for col, w in zip('ABCDE', [34, 8, 13, 26, 74]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A7'

# ─── Detalhe ─────────────────────────────────────────────────────────────────
ws2 = wb.create_sheet('Todos os 26')
cols = [('Parado (dias)', 'paradoDias', 13), ('Loja', 'loja', 32), ('Telefone', 'telefone', 16),
        ('Opp', 'opp', 9), ('Status do lead', 'statusLead', 20), ('E-mail?', 'email', 9),
        ('CNPJ', 'cnpj', 18), ('Lançado na AIVA?', 'lancado', 17),
        ('Etapa no portal', 'etapaPortal', 18), ('Último contato (dias)', 'ultimoContato', 20),
        ('Trava', 'trava', 32), ('O que fazer', 'acao', 60)]
for i, (h, _, w) in enumerate(cols, 1):
    c = ws2.cell(row=1, column=i, value=h); c.fill = HDR_BG; c.font = HDR_FT; c.border = BORDA
    ws2.column_dimensions[get_column_letter(i)].width = w
for j, l in enumerate(linhas, 2):
    for i, (_, k, _) in enumerate(cols, 1):
        v = l.get(k)
        c = ws2.cell(row=j, column=i, value=v)
        c.font = CEL; c.border = BORDA
        c.alignment = Alignment(vertical='top', wrap_text=(k == 'acao'))
        if k in ('telefone', 'cnpj'):
            c.number_format = '@'; c.value = str(v or '')
    if linhas[j - 2]['trava'].startswith('3.'):
        for i in range(1, len(cols) + 1): ws2.cell(row=j, column=i).fill = VERM
    elif (linhas[j - 2]['paradoDias'] or 0) >= 30:
        for i in range(1, len(cols) + 1):
            if not ws2.cell(row=j, column=i).fill.fgColor.rgb == 'FCE4E4': ws2.cell(row=j, column=i).fill = AMAR
ws2.freeze_panes = 'A2'
ws2.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(linhas) + 1}'

saida = 'docs/Parados-Cadastro-Recebido-2026-09-17.xlsx'
wb.save(saida)
print('gerado:', saida, f'({len(linhas)} linhas)')
