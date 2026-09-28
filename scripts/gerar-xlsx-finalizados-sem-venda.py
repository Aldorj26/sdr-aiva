# -*- coding: utf-8 -*-
"""Planilha do levantamento dos leads em LOJA_FINALIZADA_E_VENDENDO sem venda no portal (28/09/2026)."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sys.stdout.reconfigure(encoding='utf-8')

dados = json.load(open(r'C:\projetos claude\sdr-aiva\scripts\out-finalizados-sem-venda-2026-09-28.json', encoding='utf-8'))
A = 'Arial'
SUG = {
    'login': ('Login (71)', 'Loja criada e com senha, mas nunca vendeu. Etapa certa é Login: a Track liga pra destravar o uso.', 'DBEAFE'),
    'treinar': ('Treinar (70)', 'Loja criada, mas o portal não registra envio da senha. ⚠️ Mover pra 70 dispara o HSM 69 de treinamento.', 'FEF3C7'),
    'reprovado': ('Loja Descartada pela Aiva (95)', 'A AIVA reprovou este CNPJ e a loja nunca vendeu. ⚠️ O espelho não mexe em card na 51 — decisão humana.', 'FEE2E2'),
    'em_analise': ('Em Análise AIVA (50)', 'Cadastro ainda aberto no portal (formulário ou biometria). Na 50 entra na cobrança e na biometria automáticas.', 'EDE9FE'),
    'conferir': ('Conferir na mão', 'CNPJ gravado não aparece no portal da Track (ou não há CNPJ). Pode vender por outro CNPJ ou ter sido cadastrado por fora.', 'F3F4F6'),
}
wb = Workbook()
ws = wb.active; ws.title = 'Resumo'
ws.append(['Leads "Loja Finalizada e Vendendo" SEM venda no portal da AIVA — 28/09/2026']); ws['A1'].font = Font(name=A, bold=True, size=13)
ws.append(['Nada foi alterado: é a sugestão de etapa pela situação real de cada CNPJ no portal. Card hoje: todos na etapa 51.']); ws['A2'].font = Font(name=A, size=9, italic=True)
ws.append([])
ws.append(['Etapa sugerida', 'Qtd', 'Por quê / atenção'])
for c in ws[4]: c.font = Font(name=A, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F2937')
for k in ['login', 'treinar', 'reprovado', 'em_analise', 'conferir']:
    n = sum(1 for d in dados if d['sugestao'] == k)
    ws.append([SUG[k][0], n, SUG[k][1]])
    for c in ws[ws.max_row]: c.font = Font(name=A); c.fill = PatternFill('solid', fgColor=SUG[k][2]); c.alignment = Alignment(wrap_text=True, vertical='top')
ws.append(['Total', len(dados), '']); [setattr(c, 'font', Font(name=A, bold=True)) for c in ws[ws.max_row]]
for col, w in zip('ABC', (32, 8, 95)): ws.column_dimensions[col].width = w

wd = wb.create_sheet('Detalhe')
cab = ['Etapa sugerida', 'Loja', 'Telefone', 'Cidade', 'CNPJ(s)', 'Situação no portal', 'Card hoje', 'Silêncio do lojista (dias)']
wd.append(cab)
for c in wd[1]: c.font = Font(name=A, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F2937'); c.alignment = Alignment(wrap_text=True, vertical='center')
for d in dados:
    wd.append([SUG[d['sugestao']][0], d['nome'], d['telefone'], d['cidade'] or '', d['cnpjs'], d['portal'], d['card'], d['silencio'] if d['silencio'] is not None else 'nunca falou'])
    r = wd.max_row
    for c in wd[r]: c.font = Font(name=A, size=10); c.alignment = Alignment(wrap_text=True, vertical='top'); c.fill = PatternFill('solid', fgColor=SUG[d['sugestao']][2])
    wd.cell(r, 3).number_format = '@'; wd.cell(r, 5).number_format = '@'
for i, w in enumerate((26, 30, 16, 16, 36, 70, 10, 12), 1): wd.column_dimensions[get_column_letter(i)].width = w
wd.freeze_panes = 'A2'; wd.auto_filter.ref = f'A1:H{wd.max_row}'
dest = r'C:\projetos claude\sdr-aiva\docs\Finalizados-sem-venda-2026-09-28.xlsx'
wb.save(dest); print(dest)
