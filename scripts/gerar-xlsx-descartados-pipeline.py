# -*- coding: utf-8 -*-
"""Planilha dos "Descartados no pipeline" (card vermelho do painel) — Aldo 22/09/2026.

São leads DESCARTADO na nossa base cujo telefone AINDA tem card aberto no funil 15.
Cada um foi conferido por CNPJ (registro, portal AIVA, desempenho, base Odres) e por
TELEFONE (cards nos funis 15/19/17/20 e histórico de conversa).
"""
import json, io, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FONTE = r'C:\projetos claude\sdr-aiva\scripts\out-descartados-pipeline.json'
SAIDA = r'C:\projetos claude\sdr-aiva\docs\Descartados-no-pipeline-2026-09-22.xlsx'
d = json.load(io.open(FONTE, encoding='utf-8'))

ARIAL = 'Arial'
HDR = PatternFill('solid', fgColor='1F3864')
COR = {'DESCARTE OK': PatternFill('solid', fgColor='E8F5E9'),
       'CONFERIR': PatternFill('solid', fgColor='FFF8E1'),
       'NÃO DESCARTAR': PatternFill('solid', fgColor='FCE4EC')}

wb = Workbook(); wb.remove(wb.active)

def cabecalho(ws, cols, larguras):
    ws.append(cols)
    for i, c in enumerate(cols, 1):
        cel = ws.cell(row=1, column=i)
        cel.font = Font(name=ARIAL, bold=True, color='FFFFFF', size=10)
        cel.fill = HDR
        cel.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = larguras[i - 1]
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'A2'

def fmt(ws, textuais=()):
    for row in ws.iter_rows(min_row=2):
        for cel in row:
            cel.font = Font(name=ARIAL, size=10)
            cel.alignment = Alignment(vertical='top', wrap_text=(cel.column_letter in ('C', 'D')))
            if cel.column_letter in textuais:
                cel.number_format = '@'
        f = COR.get(row[0].value)
        if f:
            row[0].fill = f
            row[0].font = Font(name=ARIAL, size=10, bold=True)
    ws.auto_filter.ref = ws.dimensions

COLS = ['Recomendação', 'Loja', 'Por quê', 'Por que foi descartado', 'Telefone', 'CNPJ(s)', 'Cidade',
        'Card no funil 15', 'Mover o card?', 'Outros funis', 'No funil 19', 'RID', 'Vendas', 'Consultas',
        'Etapa no portal', 'Biometria', 'Deu os dados', 'Base Odres', 'Encerramento enviado',
        'Msgs nossas', 'Última msg do lojista', 'Silêncio (dias)', 'Disparo', 'Opp do lead']
LARG = [17, 30, 64, 34, 16, 20, 18, 24, 14, 26, 12, 10, 9, 11, 17, 13, 13, 12, 20, 12, 20, 15, 12, 12]
TEXTUAIS = ('E', 'F', 'L', 'X')
sn = lambda v: 'SIM' if v else ''

def linha(x):
    return [x['recomendacao'], x['loja'], x['motivos'], x['por_que_descartou'], x['telefone'], x['cnpjs'] or '',
            x['cidade'] or '', x['card_15'] or '(sem card)', 'MOVER' if x['card_precisa_mover'] else '',
            x['outros_funis'] or '', sn(x['no_funil_19']), x['rid'] or '', x['vendas'], x['consultas'],
            x['stage_portal'] or ('cadastro no portal' if x['no_portal'] else '(fora do portal)'),
            x['biometria'] or '', sn(x['deu_dados']), sn(x['base_odres']), x['despedida_em'] or '',
            x['nossas_mensagens'], x['ultima_msg_lojista'] or 'nunca respondeu', x['silencio_dias'],
            x['disparo'], x['opp_lead'] or '']

ordem = {'NÃO DESCARTAR': 0, 'CONFERIR': 1, 'DESCARTE OK': 2}
conta = collections.Counter(x['recomendacao'] for x in d)
n = lambda f: sum(1 for x in d if f(x))

# ── Resumo ──────────────────────────────────────────────────────────────────
ws = wb.create_sheet('Resumo')
ws.column_dimensions['A'].width = 76
ws.column_dimensions['B'].width = 12
linhas = [
    ('DESCARTADOS NO PIPELINE — CONFERÊNCIA DE 22/09/2026', ''),
    ('Leads com status DESCARTADO na nossa base que AINDA têm card aberto no funil 15.', ''),
    ('Conferidos por CNPJ (registro, portal AIVA, desempenho, base Odres) e por TELEFONE', ''),
    ('(cards nos funis 15, 19, 17 e 20 + histórico de conversa).', ''),
    ('', ''),
    ('Total analisado', len(d)),
    ('', ''),
    ('RECOMENDAÇÃO', ''),
    ('  NÃO DESCARTAR — o CRM discorda do descarte e ninguém encerrou a conversa', conta['NÃO DESCARTAR']),
    ('  CONFERIR — decisão humana antes de encerrar', conta['CONFERIR']),
    ('  DESCARTE OK — sem sinal de vida em nenhuma fonte', conta['DESCARTE OK']),
    ('', ''),
    ('A DESCOBERTA PRINCIPAL', ''),
    ('  Receberam mensagem de encerramento pelo painel ("vou encerrar nossa conversa")', n(lambda x: bool(x['despedida_em']))),
    ('  … e mesmo assim o card ficou parado numa etapa de andamento', n(lambda x: x['card_precisa_mover'])),
    ('  → nesses, o descarte foi decisão humana; o que está errado é o CARD', ''),
    ('', ''),
    ('QUANTO VALE ESSA LISTA', ''),
    ('  Chegaram a informar o CNPJ e os dados de qualificação', n(lambda x: x['deu_dados'])),
    ('  Responderam nos últimos 30 dias', n(lambda x: x['silencio_dias'] is not None and x['silencio_dias'] <= 30)),
    ('  Nunca responderam nada', n(lambda x: x['silencio_dias'] is None)),
    ('  Com loja criada na AIVA (RID) ou vendendo', n(lambda x: bool(x['rid']) or (x['vendas'] or 0) > 0)),
    ('', ''),
    ('ETAPA DO CARD NO FUNIL 15', ''),
]
for k, q in collections.Counter(x['etapa_15'] or '(sem card)' for x in d).most_common():
    linhas.append((f'  {k}', q))
linhas += [('', ''), ('SITUAÇÃO NO PORTAL DA AIVA', '')]
for k, q in collections.Counter(x['stage_portal'] or ('cadastro no portal' if x['no_portal'] else 'fora do portal') for x in d).most_common():
    linhas.append((f'  {k}', q))
linhas += [('', ''), ('MOTIVO REGISTRADO DO DESCARTE', '')]
for k, q in collections.Counter(x['por_que_descartou'] for x in d).most_common():
    linhas.append((f'  {k[:70]}', q))

for t, v in linhas:
    ws.append([t, v])
for row in ws.iter_rows(min_row=1):
    txt = str(row[0].value or '')
    negrito = txt.isupper() and len(txt) > 3
    row[0].font = Font(name=ARIAL, size=10, bold=negrito)
    row[1].font = Font(name=ARIAL, size=10, bold=negrito)
ws['A1'].font = Font(name=ARIAL, size=12, bold=True, color='1F3864')

def aba(nome, itens):
    w = wb.create_sheet(nome)
    cabecalho(w, COLS, LARG)
    for x in itens:
        w.append(linha(x))
    fmt(w, textuais=TEXTUAIS)

aba('Todos (64)', sorted(d, key=lambda y: (ordem[y['recomendacao']], y['loja'] or 'zzz')))
aba('NAO descartar', [x for x in d if x['recomendacao'] == 'NÃO DESCARTAR'])
aba('Conferir', [x for x in d if x['recomendacao'] == 'CONFERIR'])
aba('Descarte OK', [x for x in d if x['recomendacao'] == 'DESCARTE OK'])
aba('Card precisa mover', [x for x in d if x['card_precisa_mover']])

wb.save(SAIDA)
print(SAIDA)
