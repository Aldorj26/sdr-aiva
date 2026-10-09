# -*- coding: utf-8 -*-
"""Gera a planilha de conferência da comissão Ume de setembro/2026."""
import io
import json
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
DESTINO = os.path.join(RAIZ, "docs", "Conferencia-comissao-UME-2026-09.xlsx")

ESCURO = PatternFill("solid", fgColor="0F3D3E")
VERDE = PatternFill("solid", fgColor="E3F4E6")
AMARELO = PatternFill("solid", fgColor="FFF4D6")
ARIAL = "Arial"


def cab(ws, ncols, linha=1):
    for c in range(1, ncols + 1):
        cel = ws.cell(row=linha, column=c)
        cel.fill = ESCURO
        cel.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        cel.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cel.border = Border(bottom=Side(style="thin", color="7A8B8C"))
    ws.row_dimensions[linha].height = 26


def larg(ws, ws_larguras):
    for i, w in enumerate(ws_larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def corpo(ws, primeira=2, texto=(), wrap=()):
    for row in ws.iter_rows(min_row=primeira):
        for cel in row:
            cel.font = Font(name=ARIAL, size=10)
        for col in texto:
            row[col - 1].number_format = "@"
            row[col - 1].alignment = Alignment(horizontal="left")
        for col in wrap:
            row[col - 1].alignment = Alignment(horizontal="left", wrap_text=True, vertical="top")


def main():
    lojas = json.load(io.open(os.path.join(AQUI, "out-comissao-ume-2026-09.json"), encoding="utf-8"))
    rids_fcdl = set()
    with io.open(os.path.join(AQUI, "dados-fcdl-set2026.csv"), encoding="utf-8") as fh:
        next(fh)
        for l in fh:
            if l.strip():
                p = l.split(",")
                rids_fcdl.add(p[0] + "|" + p[2])
    for l in lojas:
        l["origem"] = "FCDL" if (l["rid"] + "|" + l["varejo"]) in rids_fcdl else "Outros Varejos"

    wb = Workbook()

    # ───────── Resumo ─────────
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["O que foi conferido", "Planilha / cálculo", "E-mail / declarado", "Resultado"])
    linhas = [
        ("Contratos (as duas planilhas)", "11.558", "11.558", "OK"),
        ("Originação", "R$ 5.378.181,59", "R$ 5.378.181,59", "OK"),
        ("MDR", "R$ 343.863,99", "R$ 343.863,99", "OK"),
        ("COMISSÃO TOTAL A PAGAR", "R$ 48.027,21", "R$ 48.027,21", "OK"),
        ("Nota 1 · Odres (issuer ODRES_CRED)", "R$ 3.976,02", "R$ 3.976,02", "OK"),
        ("Nota 2 · demais (não-Odres)", "R$ 44.051,19", "R$ 44.051,19", "OK"),
        ("", "", "", ""),
        ("Outros Varejos: soma das 317 lojas x cabeçalho", "R$ 45.213,53", "R$ 45.213,52", "OK (1 centavo de arredondamento)"),
        ("Outros Varejos: linha TOTAL GERAL x cabeçalho", "R$ 45.213,52", "R$ 45.213,52", "OK"),
        ("Outros Varejos: 11 subtotais de grupo", "fecham", "fecham", "OK (±R$ 0,04 de arredondamento)"),
        ("Outros Varejos: 2ª tabela (varejo x issuer)", "343 linhas, 317 varejos", "mesmo total", "OK — nenhum varejo sobra ou falta"),
        ("Outros Varejos: UME + AIVA + ODRES", "R$ 45.213,52", "R$ 45.213,52", "OK"),
        ("FCDL: soma das 311 lojas x cabeçalho", "R$ 2.813,78", "R$ 2.813,69", "OK (9 centavos de arredondamento)"),
        ("", "", "", ""),
        ("Taxa de comissão mudou de agosto p/ setembro?", "não", "", "OK — todas iguais"),
        ("Loja AIVA que vendeu no portal e não está no relatório", "1 loja", "", "ATENÇÃO — Porto Celulares"),
        ("Loja que sumiu de agosto e vendeu em setembro", "nenhuma", "", "OK"),
    ]
    for l in linhas:
        ws.append(list(l))
    cab(ws, 4)
    larg(ws, [48, 24, 22, 38])
    corpo(ws, wrap=(1, 4))
    for row in ws.iter_rows(min_row=2):
        v = row[3].value or ""
        if v.startswith("OK"):
            row[3].fill = VERDE
        elif v.startswith("ATEN"):
            row[3].fill = AMARELO
    ws.freeze_panes = "A2"

    # ───────── Achados ─────────
    ws2 = wb.create_sheet("Achados")
    ws2.append(["#", "Achado", "Impacto", "O que fazer"])
    achados = [
        (1, "Porto Celulares (retailer 6433, CNPJ 55.511.062/0001-50) vendeu 2 vezes em setembro "
            "pelo portal da AIVA (R$ 3.060,70, 19 consultas, 11 aprovados, loja Ativa desde 28/08) "
            "e NÃO tem linha em nenhuma das duas planilhas.",
         "Comissão estimada não paga: ~R$ 55,61 (1,82% da originação, que é a taxa das lojas issuer AIVA).",
         "Cobrar da Mayte a inclusão. Vale mais pelo princípio do que pelo valor: se a loja ficou de fora "
            "sendo nova, pode acontecer de novo com loja maior."),
        (2, "IVS Criciúma (retailer 2293) aparece com MDR = R$ 0,00 e comissão de R$ 131,70.",
         "A comissão está CERTA (0,60% da originação, igual a todos os outros IVS). "
            "Quem está errado é o MDR, que deveria ser ~R$ 1.097,50 (5% da originação). "
            "O MDR total do relatório está subdeclarado nesse valor.",
         "Avisar a Ume. Não afeta o que a Track recebe neste mês, mas quebra a conferência de MDR."),
        (3, "A coluna CNPJ das duas planilhas está formatada como NÚMERO, não como texto. "
            "CNPJ com 12+ dígitos aparece em notação científica (2,91808E+13) e CNPJ com zero à "
            "esquerda perde o zero.",
         "Impede cruzar o relatório com a nossa base por CNPJ sem tratar à mão. "
            "Foi por isso que a conferência deste mês cruzou por retailer ID.",
         "Pedir à Ume que formate a coluna CNPJ como texto."),
        (4, "FCDL · INSIDE_SALES: Facilita Celulares paga 6,60% do MDR e Sigma paga 6,75%. "
            "Mesma planilha, mesmo grupo, taxas diferentes.",
         "Diferença de R$ 1,71 no mês. Em agosto o grupo tinha 3 lojas com a mesma mistura (6,644%).",
         "Só confirmar com a Mayte se as taxas são mesmo por loja. Se for contrato, está certo."),
        (5, "77 lojas que comissionaram em agosto não aparecem em setembro (R$ 10.733,38 de comissão "
            "em agosto). Dessas, 66 não são lojas AIVA (são de outros varejos, onde não temos "
            "visibilidade) e 11 são AIVA sem venda nenhuma em setembro.",
         "Nenhuma delas vendeu em setembro segundo o portal. Não há comissão faltando aqui.",
         "Nada a fazer. Listadas na aba 'Sumiram' para registro."),
        (6, "As duas planilhas não puderam ser baixadas pelo Chrome — um diálogo nativo de 'salvar como' "
            "engoliu o download. Foram salvas no Drive e lidas de lá.",
         "Sem impacto na conferência, mas os anexos de setembro agora estão soltos na raiz do seu Drive.",
         "Se quiser, mover para a pasta das comissões ou apagar."),
    ]
    for a in achados:
        ws2.append(list(a))
    cab(ws2, 4)
    larg(ws2, [5, 62, 54, 54])
    corpo(ws2, wrap=(2, 3, 4))
    ws2.freeze_panes = "A2"

    # ───────── Lojas ─────────
    ws3 = wb.create_sheet("Lojas")
    ws3.append(["Planilha", "Retailer ID", "Varejo", "Grupo", "Contratos", "Originação", "MDR",
                "Comissão", "% do MDR"])
    for l in sorted(lojas, key=lambda x: -(x["comissao"] or 0)):
        ws3.append([l["origem"], l["rid"], l["varejo"], l["grupo"], l["contratos"],
                    l["originacao"], l["mdr"], l["comissao"],
                    (l["comissao"] / l["mdr"]) if l["mdr"] else None])
    cab(ws3, 9)
    larg(ws3, [15, 12, 42, 26, 10, 15, 13, 12, 10])
    corpo(ws3, texto=(2,))
    for row in ws3.iter_rows(min_row=2):
        for c in (6, 7, 8):
            row[c - 1].number_format = '#,##0.00'
        row[8].number_format = "0.00%"
    ws3.freeze_panes = "C2"
    ws3.auto_filter.ref = "A1:I%d" % ws3.max_row

    # ───────── Sumiram / Novas ─────────
    cmp_json = os.path.join(AQUI, "out-comparacao-ume-set-ago.json")
    if os.path.exists(cmp_json):
        dados = json.load(io.open(cmp_json, encoding="utf-8"))
        for titulo, chave, rot in (("Sumiram", "sumiram", "Comissão em agosto"),
                                   ("Novas em setembro", "novas", "Comissão em setembro")):
            wsx = wb.create_sheet(titulo)
            wsx.append(["Retailer ID", "Varejo", "Grupo", rot, "Vendeu em setembro no portal AIVA?"])
            for r in dados[chave]:
                wsx.append([r["rid"], r["varejo"], r.get("grupo", ""), r["comissao"], r.get("portal", "—")])
            cab(wsx, 5)
            larg(wsx, [12, 44, 26, 18, 32])
            corpo(wsx, texto=(1,))
            for row in wsx.iter_rows(min_row=2):
                row[3].number_format = '#,##0.00'
            wsx.freeze_panes = "B2"
            wsx.auto_filter.ref = "A1:E%d" % wsx.max_row
    else:
        print("aviso: %s não existe — rode antes o comparar-comissao-ume-set-vs-ago.mjs" % cmp_json)

    if not os.path.isdir(os.path.dirname(DESTINO)):
        os.makedirs(os.path.dirname(DESTINO))
    wb.save(DESTINO)
    print("salvo em", DESTINO)


if __name__ == "__main__":
    main()
